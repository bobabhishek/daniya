import os
import io
import pytest
from PIL import Image, ImageDraw
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.services.receipt_storage_service import ReceiptStorageService
from app.services.excel_export_service import ExcelExportService
from app.services.registration_service import RegistrationService
from app.services.payment_service import PaymentService
from app.config import settings

from app.firebase import memory_db

client = TestClient(app)

USER_A_TOKEN = "test_token_user_a:usera@example.com:User A"
USER_B_TOKEN = "test_token_user_b:userb@example.com:User B"
ADMIN_TOKEN = "test_token_admin_1:teamredhawkz@gmail.com:Organizer Admin"


@pytest.fixture(autouse=True)
def clean_memory_db():
    memory_db.clear()
    yield


def generate_synthetic_receipt_bytes(amount: int, utr: str = "123456789012") -> bytes:
    """Creates a synthetic UPI receipt image containing the given amount and UTR."""
    img = Image.new('RGB', (450, 600), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.rectangle([(20, 20), (430, 100)], fill=(34, 197, 94))
    draw.text((35, 40), "Payment Successful", fill=(255, 255, 255))
    draw.text((50, 140), f"Paid to: ARPITH MANOHAR", fill=(0, 0, 0))
    draw.text((50, 180), f"Amount: Rs. {amount}", fill=(0, 0, 0))
    draw.text((50, 220), f"UPI Ref: {utr}", fill=(0, 0, 0))
    draw.text((50, 260), "Taal Pe Nacho Re 2026", fill=(50, 50, 50))
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    return buf.getvalue()


def test_01_verified_payment_stores_receipt_locally():
    """
    Test 1: When 3-way match passes (expected == entered == ocr),
    the verified receipt is stored at receipts/{registrationId}/payment_receipt.jpg
    and Firestore metadata includes receiptPath.
    """
    create_res = client.post(
        "/api/registrations",
        json={"participants": [{"name": "Pooja Verma", "dob": "10/05/2001"}]},
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    assert create_res.status_code == 201
    reg_id = create_res.json()["registrationId"]

    receipt_bytes = generate_synthetic_receipt_bytes(299)
    files = {"receipt": ("receipt.jpg", receipt_bytes, "image/jpeg")}
    data = {"registration_id": reg_id, "entered_amount": "299"}

    verify_res = client.post(
        "/api/payments/verify-proof",
        data=data,
        files=files,
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    assert verify_res.status_code == 200
    res_json = verify_res.json()
    assert res_json["success"] is True
    assert res_json["verificationStatus"] == "VERIFIED"
    assert res_json["paymentStatus"] == "PAID"
    assert res_json["receiptPath"] == f"receipts/{reg_id}/payment_receipt.jpg"

    # Confirm receipt file physically exists on disk
    expected_path = ReceiptStorageService.get_receipt_path(reg_id)
    assert expected_path is not None
    assert os.path.isfile(expected_path)
    assert os.path.getsize(expected_path) > 0


def test_02_failed_payment_does_not_create_receipt():
    """
    Test 2: When entered amount mismatches expected amount,
    payment is rejected, tickets are NOT generated, and NO receipt is retained.
    """
    create_res = client.post(
        "/api/registrations",
        json={"participants": [{"name": "Rohan Gupta", "dob": "12/12/1999"}]},
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    reg_id = create_res.json()["registrationId"]

    receipt_bytes = generate_synthetic_receipt_bytes(299)
    files = {"receipt": ("receipt.jpg", receipt_bytes, "image/jpeg")}
    # User entered 199 when expected is 299
    data = {"registration_id": reg_id, "entered_amount": "199"}

    verify_res = client.post(
        "/api/payments/verify-proof",
        data=data,
        files=files,
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    assert verify_res.status_code == 200
    res_json = verify_res.json()
    assert res_json["success"] is False
    assert res_json["verificationStatus"] == "FAILED"
    assert len(res_json["ticketIds"]) == 0

    # Ensure no verified receipt file exists on disk
    receipt_path = ReceiptStorageService.get_receipt_path(reg_id)
    assert receipt_path is None


def test_03_ocr_failure_does_not_create_receipt():
    """
    Test 3: When OCR cannot read the screenshot (e.g. blank image),
    verification fails, tickets are NOT generated, and NO receipt is retained.
    """
    create_res = client.post(
        "/api/registrations",
        json={"participants": [{"name": "Kavita Rao", "dob": "15/07/2002"}]},
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    reg_id = create_res.json()["registrationId"]

    # Plain blank image with no text
    blank_img = Image.new('RGB', (200, 200), color=(250, 250, 250))
    buf = io.BytesIO()
    blank_img.save(buf, format="JPEG")

    files = {"receipt": ("blank.jpg", buf.getvalue(), "image/jpeg")}
    data = {"registration_id": reg_id, "entered_amount": "299"}

    verify_res = client.post(
        "/api/payments/verify-proof",
        data=data,
        files=files,
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    assert verify_res.status_code == 200
    res_json = verify_res.json()
    assert res_json["success"] is False
    assert res_json["verificationStatus"] == "FAILED"

    # Ensure no verified receipt file exists on disk
    receipt_path = ReceiptStorageService.get_receipt_path(reg_id)
    assert receipt_path is None


def test_04_master_excel_export_succeeds_and_returns_xlsx():
    """
    Test 4: Admin can download Master Excel via GET /api/admin/export-excel
    and response is a valid .xlsx spreadsheet.
    """
    res = client.get(
        "/api/admin/export-excel",
        headers={"Authorization": f"Bearer {ADMIN_TOKEN}"}
    )
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    assert "Dandiya_Master_Registrations" in res.headers["content-disposition"]
    assert len(res.content) > 1000


def test_05_master_excel_contains_physically_embedded_images():
    """
    Test 5: Master Excel physically embeds payment screenshot images in Column G
    for verified registrations, and is self-contained without needing internet or external URLs.
    """
    # 1. Create and verify a booking with an image
    create_res = client.post(
        "/api/registrations",
        json={"participants": [{"name": "Vikram Singh", "dob": "20/03/1995"}]},
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    reg_id = create_res.json()["registrationId"]

    receipt_bytes = generate_synthetic_receipt_bytes(299)
    client.post(
        "/api/payments/verify-proof",
        data={"registration_id": reg_id, "entered_amount": "299"},
        files={"receipt": ("receipt.jpg", receipt_bytes, "image/jpeg")},
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )

    # 2. Download Master Excel workbook
    res = client.get(
        "/api/admin/export-excel",
        headers={"Authorization": f"Bearer {ADMIN_TOKEN}"}
    )
    assert res.status_code == 200

    # 3. Load with openpyxl and inspect embedded drawings
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb["Master Registrations"]

    # Verify headers
    headers = [cell.value for cell in ws[1]]
    assert "Registration ID" in headers
    assert "Payment Screenshot" in headers
    assert "Amount" in headers
    assert "Ticket Link(s)" in headers

    # Verify that images are embedded in the worksheet
    assert hasattr(ws, "_images")
    assert len(ws._images) >= 1

    # Verify image dimensions
    embedded_img = ws._images[0]
    assert embedded_img.width > 0
    assert embedded_img.height > 0
    assert embedded_img.width <= 250


def test_06_master_excel_multiple_registrations_correct_rows():
    """
    Test 6: Multiple registrations place images in the appropriate rows,
    while unverified rows have no image in Column G.
    """
    # Create one verified and one pending registration
    c1 = client.post(
        "/api/registrations",
        json={"participants": [{"name": "Participant One", "dob": "01/01/1990"}]},
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    reg_id_1 = c1.json()["registrationId"]

    # Verified
    r_bytes = generate_synthetic_receipt_bytes(299)
    client.post(
        "/api/payments/verify-proof",
        data={"registration_id": reg_id_1, "entered_amount": "299"},
        files={"receipt": ("r.jpg", r_bytes, "image/jpeg")},
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )

    # Pending registration (no payment submitted)
    c2 = client.post(
        "/api/registrations",
        json={"participants": [{"name": "Participant Two", "dob": "02/02/1992"}]},
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    reg_id_2 = c2.json()["registrationId"]

    # Export Excel
    res = client.get(
        "/api/admin/export-excel",
        headers={"Authorization": f"Bearer {ADMIN_TOKEN}"}
    )
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb["Master Registrations"]

    # Find row index for reg_id_1 and reg_id_2
    row_1_idx = None
    row_2_idx = None
    for r in range(2, ws.max_row + 1):
        val = ws.cell(row=r, column=1).value
        if val == reg_id_1:
            row_1_idx = r
        elif val == reg_id_2:
            row_2_idx = r

    assert row_1_idx is not None
    assert row_2_idx is not None

    # Check row heights: verified row is taller to accommodate the thumbnail
    assert ws.row_dimensions[row_1_idx].height > 40
    # Pending row has standard height
    assert ws.row_dimensions[row_2_idx].height <= 30


def test_07_excel_export_still_works_if_receipt_file_is_missing():
    """
    Test 7: If a registration is marked verified but its physical file is missing,
    Excel export still succeeds gracefully without crashing.
    """
    # Mark a mock registration with non-existent receipt
    RegistrationService.mark_payment_verified(
        registration_id="KD-GHOST-999",
        entered_amount=299,
        ocr_amount=299,
        receipt_path="receipts/KD-GHOST-999/non_existent.jpg"
    )

    res = client.get(
        "/api/admin/export-excel",
        headers={"Authorization": f"Bearer {ADMIN_TOKEN}"}
    )
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))
    assert "Master Registrations" in wb.sheetnames


def test_08_no_google_drive_routes_exist():
    """
    Test 8: Confirms that Google Drive routes are completely removed and return 404.
    """
    res1 = client.get("/api/drive/auth")
    assert res1.status_code == 404

    res2 = client.get("/api/drive/status")
    assert res2.status_code == 404

    res3 = client.get("/api/drive/oauth2callback")
    assert res3.status_code == 404


def test_09_master_excel_contains_participant_name_dob_and_phone():
    """
    Test 9: Verifies that when participant details including phone number are registered,
    the generated Excel workbook contains Name, Date of Birth, and Phone Number
    in both the Master Registrations sheet and the Participant Details sheet.
    """
    create_res = client.post(
        "/api/registrations",
        json={
            "participants": [
                {
                    "name": "Meera Patel",
                    "dob": "14/05/2001",
                    "phone": "9876543210"
                }
            ]
        },
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    assert create_res.status_code == 201
    reg_id = create_res.json()["registrationId"]
    assert create_res.json()["participants"][0]["phone"] == "9876543210"

    res = client.get(
        "/api/admin/export-excel",
        headers={"Authorization": f"Bearer {ADMIN_TOKEN}"}
    )
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))

    # 1. Master Registrations Sheet: Column C contains Name, DOB, and Phone
    ws_master = wb["Master Registrations"]
    found_row = None
    for r in range(2, ws_master.max_row + 1):
        if ws_master.cell(row=r, column=1).value == reg_id:
            found_row = r
            break
    assert found_row is not None
    summary_val = str(ws_master.cell(row=found_row, column=3).value)
    assert "Meera Patel" in summary_val
    assert "14/05/2001" in summary_val
    assert "9876543210" in summary_val

    # 2. Participant Details Sheet: dedicated columns for Name, DOB, Phone
    assert "Participant Details" in wb.sheetnames
    ws_parts = wb["Participant Details"]
    headers = [cell.value for cell in ws_parts[1]]
    assert "Full Name" in headers
    assert "Date of Birth" in headers
    assert "Phone Number" in headers

    # Locate participant row
    part_row_found = False
    for r in range(2, ws_parts.max_row + 1):
        if ws_parts.cell(row=r, column=1).value == reg_id:
            assert ws_parts.cell(row=r, column=3).value == "Meera Patel"
            assert ws_parts.cell(row=r, column=4).value == "14/05/2001"
            assert ws_parts.cell(row=r, column=5).value == "9876543210"
            part_row_found = True
            break
    assert part_row_found is True


def test_10_phone_validation_in_registration_api():
    """
    Test 10: Verifies that Indian 10-digit mobile number format is validated.
    - Valid formats with +91 or leading 0 are cleaned to 10 digits.
    - Invalid formats (wrong length or starting digit) return 422 Unprocessable Entity.
    """
    # 1. Phone with +91 cleaned to 10 digits
    res_clean = client.post(
        "/api/registrations",
        json={"participants": [{"name": "Aarav Shah", "dob": "01/01/2000", "phone": "+91 91234 56789"}]},
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    assert res_clean.status_code == 201
    assert res_clean.json()["participants"][0]["phone"] == "9123456789"

    # 2. Invalid phone (too short) returns 422
    res_invalid_short = client.post(
        "/api/registrations",
        json={"participants": [{"name": "Aarav Shah", "dob": "01/01/2000", "phone": "98765"}]},
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    assert res_invalid_short.status_code == 422

    # 3. Invalid phone (starts with 3, not valid Indian mobile) returns 422
    res_invalid_digit = client.post(
        "/api/registrations",
        json={"participants": [{"name": "Aarav Shah", "dob": "01/01/2000", "phone": "3123456789"}]},
        headers={"Authorization": f"Bearer {USER_A_TOKEN}"}
    )
    assert res_invalid_digit.status_code == 422
